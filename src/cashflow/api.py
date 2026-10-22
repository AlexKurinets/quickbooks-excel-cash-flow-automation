"""HTTP API for the dashboard and for anything else the firm wants to plug in (a Teams bot, a portal)."""

from __future__ import annotations

import calendar
import hashlib
import io
import zipfile
from datetime import UTC, date, datetime

import openpyxl
from fastapi import FastAPI, HTTPException, Query, Response
from pydantic import BaseModel

from . import db, mapping, runner
from .config import settings
from .lines import LINES
from .synthetic import STAFF
from .workbook import TemplateError, TemplateMap, _full_calc_on_load, formula_inventory, locate, sheet_parts

app = FastAPI(title="QuickBooks to Excel cash flow automation", version="1.0.0")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _period(text: str) -> date:
    try:
        year, month = (int(x) for x in text.split("-")[:2])
        return date(year, month, calendar.monthrange(year, month)[1])
    except Exception as exc:
        raise HTTPException(422, f"period must look like 2026-09, got {text!r}") from exc


class Actor(BaseModel):
    actor: str


class RunRequest(Actor):
    client_id: str
    period: str


class Approve(Actor):
    line_key: str | None = None


class Reject(Actor):
    note: str


@app.get("/health")
def health() -> dict:
    with db.connect() as conn:
        conn.execute("select 1")
    return {"ok": True, "source": settings().source}


@app.get("/staff")
def staff() -> list[dict]:
    return [{"email": e, "name": n, "role": r} for e, n, r in STAFF]


@app.get("/lines")
def lines() -> list[dict]:
    return [{"key": line.key, "section": line.section, "title": line.title, "description": line.description}
            for line in LINES]


@app.get("/periods")
def periods() -> list[str]:
    with db.connect() as conn:
        rows = conn.execute("select distinct period from runs order by period desc").fetchall()
    return [f"{r['period']:%Y-%m}" for r in rows]


@app.get("/clients")
def clients() -> list[dict]:
    with db.connect() as conn:
        return conn.execute("select id, name, industry, realm_id, connection_status, connection_expired_at "
                            "from clients order by name").fetchall()


@app.get("/close")
def close(period: str = Query(...)) -> dict:
    p = _period(period)
    with db.connect() as conn:
        rows = conn.execute(
            """select * from (
                 select distinct on (c.id) c.id as client_id, c.name, c.industry, c.connection_status,
                        r.id as run_id, r.status, r.reason, r.started_at, r.duration_ms, r.delivered_at,
                        r.delivered_by, r.created_by,
                        (select count(*) from run_checks k where k.run_id = r.id and k.status = 'warn') as warnings,
                        (select summary from run_checks k where k.run_id = r.id and k.status = 'warn'
                         order by ordinal limit 1) as warning
                 from clients c left join runs r on r.client_id = c.id and r.period = %s
                 order by c.id, r.id desc) latest
               order by name""", (p,)).fetchall()
        pending = conn.execute("select count(*) as n from mapping_suggestions where status = 'pending'").fetchone()["n"]
    counts = {s: sum(1 for r in rows if r["status"] == s) for s in ("ready", "held", "failed")}
    counts["sent"] = sum(1 for r in rows if r["delivered_at"])
    counts["not_run"] = sum(1 for r in rows if r["status"] is None)
    return {"period": f"{p:%Y-%m}", "counts": counts, "pending_mappings": pending, "clients": rows}


def _run(conn, run_id: int) -> dict:
    run = conn.execute(
        """select r.id, r.client_id, c.name as client_name, c.industry, c.realm_id, r.period, r.as_of, r.status,
                  r.reason, r.started_at, r.finished_at, r.duration_ms, r.template_name, r.template_sha256,
                  r.formula_count, r.workbook_sha256, length(r.workbook) as workbook_bytes, r.figures, r.created_by,
                  r.delivered_at, r.delivered_by
           from runs r join clients c on c.id = r.client_id where r.id = %s""", (run_id,)).fetchone()
    if run is None:
        raise HTTPException(404, f"no run {run_id}")
    run["checks"] = conn.execute("select key, title, status, blocking, summary, detail from run_checks "
                                 "where run_id = %s order by ordinal", (run_id,)).fetchall()
    run["period"] = f"{run['period']:%Y-%m}"
    return run


@app.get("/runs")
def runs(client_id: str | None = None, limit: int = 50) -> list[dict]:
    with db.connect() as conn:
        return conn.execute(
            """select r.id, r.client_id, c.name as client_name, to_char(r.period, 'YYYY-MM') as period, r.status,
                      r.reason, r.started_at, r.duration_ms, r.delivered_at, r.created_by
               from runs r join clients c on c.id = r.client_id
               where %s::text is null or r.client_id = %s order by r.id desc limit %s""",
            (client_id, client_id, limit)).fetchall()


@app.get("/runs/{run_id}")
def run_detail(run_id: int) -> dict:
    with db.connect() as conn:
        return _run(conn, run_id)


@app.post("/runs")
def start_run(req: RunRequest) -> dict:
    with db.connect() as conn:
        if conn.execute("select 1 from clients where id = %s", (req.client_id,)).fetchone() is None:
            raise HTTPException(404, f"no client {req.client_id}")
        run_id = runner.run_client(conn, req.client_id, _period(req.period), actor=req.actor)
        return _run(conn, run_id)


@app.post("/runs/{run_id}/deliver")
def deliver(run_id: int, req: Actor) -> dict:
    with db.connect() as conn:
        try:
            runner.deliver(conn, run_id, req.actor)
        except KeyError as exc:
            raise HTTPException(404, str(exc)) from exc
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from exc
        conn.commit()
        return _run(conn, run_id)


def _workbook_response(row: dict) -> Response:
    if not row["workbook"]:
        raise HTTPException(409, f"run {row['id']} is {row['status']} and has no workbook: {row['reason']}")
    name = f"{row['client_name']} - Cash Flow - {row['period']:%Y-%m}.xlsx".replace("&", "and")
    return Response(bytes(row["workbook"]), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{name}"', "X-Run-Id": str(row["id"]),
                             "X-Run-Status": row["status"]})


@app.get("/runs/{run_id}/workbook")
def workbook(run_id: int) -> Response:
    with db.connect() as conn:
        row = conn.execute("select r.*, c.name as client_name from runs r join clients c on c.id = r.client_id "
                           "where r.id = %s", (run_id,)).fetchone()
    if row is None:
        raise HTTPException(404, f"no run {run_id}")
    return _workbook_response(row)


@app.get("/clients/{client_id}/workbook")
def client_workbook(client_id: str, period: str) -> Response:
    with db.connect() as conn:
        row = conn.execute("select r.*, c.name as client_name from runs r join clients c on c.id = r.client_id "
                           "where r.client_id = %s and r.period = %s order by r.id desc limit 1",
                           (client_id, _period(period))).fetchone()
    if row is None:
        raise HTTPException(404, f"no run for {client_id} in {period}")
    return _workbook_response(row)


@app.get("/runs/{run_id}/preview")
def preview(run_id: int) -> dict:
    """The Cash Flow sheet as any viewer that doesn't recalculate would show it: cached values only."""
    with db.connect() as conn:
        row = conn.execute("select workbook, status from runs where id = %s", (run_id,)).fetchone()
    if row is None:
        raise HTTPException(404, f"no run {run_id}")
    if not row["workbook"]:
        raise HTTPException(409, f"run {run_id} is {row['status']} and has no workbook")
    s = settings()
    tmap = TemplateMap.load(s.template_map)
    wb = openpyxl.load_workbook(io.BytesIO(bytes(row["workbook"])), data_only=True)
    ws = wb[tmap.sheet]
    header = None
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=14):
        label = r[0].value
        if label == tmap.month_header:
            header = [c.value for c in r[1:]]
            continue
        if header is None or not label:
            continue
        rows.append({"label": label, "values": [c.value for c in r[1:]], "bold": bool(r[0].font and r[0].font.b)})
    fields = {k: ws[f"B{i}"].value for i, k in enumerate(["client", "fiscal_year", "closed_through", "source"], 2)}
    return {"fields": fields, "columns": header, "rows": rows}


@app.get("/suggestions")
def suggestions(status: str = "pending") -> list[dict]:
    with db.connect() as conn:
        if status == "pending":
            where = "s.status = 'pending'"
        else:
            where = "s.status <> 'pending'"
        return conn.execute(
            f"""select s.*, c.name as client_name, c.industry from mapping_suggestions s
                join clients c on c.id = s.client_id where {where}
                order by coalesce(s.decided_at, s.created_at) desc, s.id desc limit 50""").fetchall()


@app.post("/suggestions/{suggestion_id}/approve")
def approve(suggestion_id: int, req: Approve) -> dict:
    with db.connect() as conn:
        try:
            decided, run_id = runner.approve_and_rerun(conn, suggestion_id, req.actor, req.line_key)
        except mapping.MappingError as exc:
            raise HTTPException(409, str(exc)) from exc
        return {"suggestion": decided, "run": _run(conn, run_id) if run_id else None}


@app.post("/suggestions/{suggestion_id}/reject")
def reject(suggestion_id: int, req: Reject) -> dict:
    with db.connect() as conn:
        try:
            return mapping.reject(conn, suggestion_id, req.actor, datetime.now(UTC), req.note)
        except mapping.MappingError as exc:
            raise HTTPException(409, str(exc)) from exc


@app.post("/clients/{client_id}/reconnect")
def reconnect(client_id: str, req: Actor) -> dict:
    if settings().source != "synthetic":
        raise HTTPException(409, "Reconnect through Intuit's consent screen; this shortcut exists only for synthetic data.")
    with db.connect() as conn:
        runner.reconnect(conn, client_id, req.actor)
    return {"client_id": client_id, "connection_status": "connected"}


@app.get("/template")
def template() -> dict:
    s = settings()
    data = s.template.read_bytes()
    tmap = TemplateMap.load(s.template_map)
    zf = zipfile.ZipFile(io.BytesIO(data))
    sheets = sheet_parts(zf)
    try:
        layout = locate(data, tmap)
        problem = None
    except TemplateError as exc:
        layout, problem = None, str(exc)
    inputs = [{"key": k, "label": spec.label, "row": layout.rows[k] if layout else None,
               "months": "first month only" if spec.first_month_only else "every month"}
              for k, spec in tmap.inputs.items()]
    outputs = [{"key": k, "label": label, "row": layout.rows[k] if layout else None} for k, label in tmap.outputs.items()]
    inventory = formula_inventory(data)
    by_sheet = {name: sum(1 for k in inventory if k.startswith(f"{name}!")) for name in sheets}
    written = {sheets[tmap.sheet]} | {sheets[name] for name, count in by_sheet.items() if count}
    flag_set = _full_calc_on_load(zf.read("xl/workbook.xml")) is None

    def writer(part: str) -> str:
        if part in written:
            return "input cells and formula results only"
        if part == "xl/workbook.xml" and not flag_set:
            return "adds the recalculate-on-open flag"
        return "never touched"

    parts = [{"part": info.filename, "bytes": info.file_size, "writer": writer(info.filename)} for info in zf.infolist()]
    return {"name": s.template.name, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data),
            "formula_count": len(inventory), "formulas_by_sheet": by_sheet, "sheet": tmap.sheet,
            "fields": [{"key": k, "label": label, "cell": layout.fields[k] if layout else None}
                       for k, label in tmap.fields.items()],
            "inputs": inputs, "outputs": outputs, "parts": parts, "problem": problem}
