"""The seeded closes end to end. Read-only assertions first; the tests at the bottom change state."""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, date, datetime
from pathlib import Path

import openpyxl
from lxml import etree

from cashflow import runner
from cashflow.workbook import Q, sheet_parts
from conftest import TEMPLATE, checks, latest

LATER = datetime(2026, 10, 12, 15, 0, tzinfo=UTC)
READY = {"wrenfield-dental", "tamsin-creek-brewing", "ollery-lane-property", "northvane-mechanical", "quillmark-studio"}
HELD = {"birchline-grounds": "mapping_complete", "saltbrook-animal-clinic": "equity_rollforward",
        "rye-and-ember-bakery": "qbo_connection"}


def test_august_close_was_clean_and_sent(seeded):
    conn = seeded["conn"]
    rows = conn.execute("select status, delivered_at from runs where period = '2026-08-31'").fetchall()
    assert len(rows) == 8 and all(r["status"] == "ready" and r["delivered_at"] for r in rows)


def test_september_close_holds_exactly_the_three_clients_with_a_real_problem(seeded):
    conn = seeded["conn"]
    for client_id in READY:
        assert latest(conn, client_id)["status"] == "ready", client_id
    for client_id, key in HELD.items():
        run = latest(conn, client_id)
        assert run["status"] == "held", client_id
        failing = [k for k, c in checks(conn, run["id"]).items() if c["status"] == "fail"]
        assert failing[0] == key
        assert run["reason"] == checks(conn, run["id"])[key]["summary"]


def test_new_account_files_a_suggestion_and_builds_no_workbook(seeded):
    conn = seeded["conn"]
    run = latest(conn, "birchline-grounds")
    assert run["workbook"] is None
    s = conn.execute("select * from mapping_suggestions where status = 'pending'").fetchall()
    assert [(x["client_id"], x["account_name"], x["suggested_line"], x["run_id"]) for x in s] == [
        ("birchline-grounds", "Equipment Loan - Mower Fleet", "loans", run["id"])]
    assert "Equipment Loan - Mower Fleet" in run["reason"]


def test_retained_earnings_posting_is_explained_and_every_cash_check_agrees(seeded):
    conn = seeded["conn"]
    c = checks(conn, latest(conn, "saltbrook-animal-clinic")["id"])
    assert "In September equity moved $2,350.00 less than net income" in c["equity_rollforward"]["summary"]
    for key in ("ending_cash_ties", "cash_flow_report_ties", "template_difference_zero"):
        assert c[key]["status"] == "fail"
    assert "off by $2,350.00" in c["ending_cash_ties"]["summary"]
    assert c["formulas_preserved"]["status"] == "pass"


def test_a_month_that_changed_after_it_was_sent_is_a_warning_not_a_hold(seeded):
    conn = seeded["conn"]
    run = latest(conn, "northvane-mechanical")
    w = checks(conn, run["id"])["delivered_months_unchanged"]
    assert run["status"] == "ready" and w["status"] == "warn" and not w["blocking"]
    assert w["summary"].startswith("August changed in QuickBooks after the August 2026 workbook was sent on 9 Sep 2026")
    assert "$1,240.00 lower" in w["summary"]
    assert w["detail"]["months"] == [{"month": "Aug", "cash_change": -1240.0, "net_income_change": -1240.0,
                                      "cash_then": w["detail"]["months"][0]["cash_then"],
                                      "cash_now": w["detail"]["months"][0]["cash_now"]}]


def test_ready_workbook_ties_and_its_untouched_parts_match_the_template(seeded):
    conn = seeded["conn"]
    run = latest(conn, "wrenfield-dental")
    book = bytes(run["workbook"])
    tpl = Path(TEMPLATE).read_bytes()
    zt, zb = zipfile.ZipFile(io.BytesIO(tpl)), zipfile.ZipFile(io.BytesIO(book))
    for name in ("xl/charts/chart1.xml", "xl/drawings/drawing1.xml", "xl/styles.xml", "xl/workbook.xml"):
        assert zt.read(name) == zb.read(name)
    ws = openpyxl.load_workbook(io.BytesIO(book), data_only=True)["Cash Flow"]
    figures = run["figures"]
    assert ws["J37"].value == figures["qbo"]["bs_cash"][8] == figures["qbo"]["cf_ending_cash"][8]
    assert ws["B4"].value == "September 2026"
    assert ws["B5"].value.startswith("QuickBooks Online company 4620816365000102, data as of 2026-10-08 09:00 UTC")
    assert all(c["status"] == "pass" for c in checks(conn, run["id"]).values())


def test_approving_the_suggestion_reruns_birchline_to_ready(seeded):
    conn = seeded["conn"]
    sid = conn.execute("select id from mapping_suggestions where status = 'pending'").fetchone()["id"]
    decided, run_id = runner.approve_and_rerun(conn, sid, "mara.quist@example.com", now=LATER)
    assert decided["decided_line"] == "loans"
    run = conn.execute("select * from runs where id = %s", (run_id,)).fetchone()
    assert run["status"] == "ready" and run["period"] == date(2026, 9, 30)
    sep = run["figures"]["workbook"]["ending_cash"][8]
    assert sep == run["figures"]["qbo"]["bs_cash"][8]
    loans = run["figures"]["lines"]["loans"][8]
    assert loans > 40_000  # the $48,000 loan less September's payment on the older loan
    assert run["figures"]["lines"]["property_equipment"][8] == -56_500.0


def test_reconnecting_lets_the_bakery_run(seeded):
    conn = seeded["conn"]
    runner.reconnect(conn, "rye-and-ember-bakery", "mara.quist@example.com", LATER)
    run_id = runner.run_client(conn, "rye-and-ember-bakery", date(2026, 9, 30), actor="mara.quist@example.com",
                               now=LATER)
    assert conn.execute("select status from runs where id = %s", (run_id,)).fetchone()["status"] == "ready"


def test_a_template_someone_broke_holds_the_run_and_names_the_cell(seeded, tmp_path):
    conn = seeded["conn"]
    tpl = Path(TEMPLATE).read_bytes()
    zf = zipfile.ZipFile(io.BytesIO(tpl))
    path = sheet_parts(zf)["Cash Flow"]
    root = etree.fromstring(zf.read(path))
    cell = next(c for c in root.iter(f"{Q}c") if c.get("r") == "J35")
    cell.remove(cell.find(f"{Q}f"))
    broken = tmp_path / "broken.xlsx"
    with zipfile.ZipFile(broken, "w") as z:
        for info in zf.infolist():
            z.writestr(info, etree.tostring(root) if info.filename == path else zf.read(info.filename))
    run_id = runner.run_client(conn, "wrenfield-dental", date(2026, 9, 30), now=LATER, template_path=broken)
    run = conn.execute("select * from runs where id = %s", (run_id,)).fetchone()
    assert run["status"] == "held" and run["workbook"] is None
    assert "J35" in run["reason"] and "typed over" in run["reason"]


def test_only_ready_runs_can_be_marked_sent(seeded):
    conn = seeded["conn"]
    held = latest(conn, "saltbrook-animal-clinic")
    try:
        runner.deliver(conn, held["id"], "mara.quist@example.com")
        raise AssertionError("held run was delivered")
    except ValueError as exc:
        assert "only a ready workbook" in str(exc)
    conn.rollback()
