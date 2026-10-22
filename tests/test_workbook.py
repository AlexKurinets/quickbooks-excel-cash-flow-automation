"""The writer's guarantees: inputs only, formulas and every other part untouched, real cached values."""

from __future__ import annotations

import io
import sys
import zipfile
from decimal import Decimal

import openpyxl
import pytest
from lxml import etree

from cashflow.workbook import Q, TemplateError, TemplateMap, fill, formula_inventory, locate, patch, sheet_parts
from conftest import HOME, MAP

sys.path.insert(0, str(HOME / "scripts"))
import build_template  # noqa: E402

TMAP = TemplateMap.load(MAP)
DISCLAIMER = "without audit or review"


def sample_inputs(layout, months=9):
    inputs = {layout.fields["client_name"]: "Example Co", layout.fields["fiscal_year"]: 2026,
              layout.cell("cash_beginning", "Jan"): Decimal("10000.00")}
    cash = Decimal("10000.00")
    for i, m in enumerate(TMAP.months[:months]):
        ni, ar, dist = Decimal("1500.25") + i, Decimal("-300.10"), Decimal("-500.00")
        inputs[layout.cell("net_income", m)] = ni
        inputs[layout.cell("accounts_receivable", m)] = ar
        inputs[layout.cell("owner_distributions", m)] = dist
        cash += ni + ar + dist
        inputs[layout.cell("cash_per_balance_sheet", m)] = cash
    return inputs, cash


def part(xlsx: bytes, name: str) -> bytes:
    return zipfile.ZipFile(io.BytesIO(xlsx)).read(name)


def test_locate_finds_every_target_by_label(template_bytes):
    layout = locate(template_bytes, TMAP)
    assert layout.columns["Jan"] == "B" and layout.columns["Dec"] == "M"
    assert layout.rows["net_income"] == 9 and layout.rows["ending_cash"] == 37
    assert layout.fields == {"client_name": "B2", "fiscal_year": "B3", "closed_through": "B4", "source": "B5"}


def test_fill_writes_inputs_only_and_leaves_formulas_chart_and_text_box_byte_identical(template_bytes):
    layout = locate(template_bytes, TMAP)
    inputs, cash = sample_inputs(layout)
    filled = fill(template_bytes, layout, inputs)
    v = filled.verification
    assert v.ok and v.formula_count == 172
    assert formula_inventory(filled.workbook) == formula_inventory(template_bytes)
    assert set(v.changed_parts) == {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"}
    for name in ("xl/charts/chart1.xml", "xl/drawings/drawing1.xml", "xl/styles.xml", "xl/sharedStrings.xml",
                 "xl/workbook.xml", "xl/worksheets/sheet3.xml"):
        assert part(filled.workbook, name) == part(template_bytes, name)
    assert DISCLAIMER in part(filled.workbook, "xl/drawings/drawing1.xml").decode()
    assert filled.value("ending_cash", "Sep") == pytest.approx(float(cash))
    assert filled.value("difference", "Sep") == 0
    assert filled.value("difference", "Oct") == ""


def test_cached_values_are_what_a_non_recalculating_viewer_reads(template_bytes):
    layout = locate(template_bytes, TMAP)
    inputs, cash = sample_inputs(layout)
    filled = fill(template_bytes, layout, inputs)
    wb = openpyxl.load_workbook(io.BytesIO(filled.workbook), data_only=True)
    ws, summary = wb["Cash Flow"], wb["Summary"]
    assert ws[layout.cell("ending_cash", "Sep")].value == pytest.approx(float(cash))
    assert summary["A2"].value == "Example Co"
    assert summary["A3"].value == "Fiscal year 2026"
    assert summary["F14"].value == pytest.approx(float(cash))  # September ending cash on the summary sheet
    assert summary["F15"].value == "#N/A"                       # October not filled yet, so the chart skips it


def test_writing_into_a_formula_cell_is_refused(template_bytes):
    layout = locate(template_bytes, TMAP)
    with pytest.raises(TemplateError, match="formula"):
        patch(template_bytes, {"Cash Flow": {layout.cell("ending_cash", "Mar"): 1}})


def test_a_duplicated_label_stops_the_run_before_anything_is_written(template_bytes):
    broken = patch(template_bytes, {"Cash Flow": {"A50": "Net income"}})
    with pytest.raises(TemplateError, match="appears 2 times"):
        locate(broken, TMAP)


def test_a_subtotal_someone_typed_over_is_caught(template_bytes):
    zf = zipfile.ZipFile(io.BytesIO(template_bytes))
    path = sheet_parts(zf)["Cash Flow"]
    root = etree.fromstring(zf.read(path))
    cell = next(c for c in root.iter(f"{Q}c") if c.get("r") == "E21")
    cell.remove(cell.find(f"{Q}f"))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for info in zf.infolist():
            z.writestr(info, etree.tostring(root) if info.filename == path else zf.read(info.filename))
    with pytest.raises(TemplateError, match="E21.*should be a formula"):
        locate(out.getvalue(), TMAP)


def test_an_inserted_row_moves_the_numbers_with_it(tmp_path, template_bytes):
    moved = build_template.build(tmp_path / "moved.xlsx", extra_row=True).read_bytes()
    a, b = locate(template_bytes, TMAP), locate(moved, TMAP)
    assert b.rows["depreciation"] == a.rows["depreciation"] + 1
    assert b.rows["ending_cash"] == a.rows["ending_cash"] + 1
    fa = fill(template_bytes, a, sample_inputs(a)[0])
    fb = fill(moved, b, sample_inputs(b)[0])
    assert [fa.value("ending_cash", m) for m in TMAP.months] == [fb.value("ending_cash", m) for m in TMAP.months]
    assert fb.verification.ok


def test_recalculate_on_open_flag_is_added_when_missing(template_bytes):
    zf = zipfile.ZipFile(io.BytesIO(template_bytes))
    root = etree.fromstring(zf.read("xl/workbook.xml"))
    root.find(f"{Q}calcPr").attrib.pop("fullCalcOnLoad")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for info in zf.infolist():
            z.writestr(info, etree.tostring(root) if info.filename == "xl/workbook.xml" else zf.read(info.filename))
    patched = patch(out.getvalue(), {"Cash Flow": {"B2": "Example Co"}})
    calc = etree.fromstring(part(patched, "xl/workbook.xml")).find(f"{Q}calcPr")
    assert calc.get("fullCalcOnLoad") == "1"


def test_why_the_writer_does_not_resave_with_openpyxl(template_bytes, tmp_path):
    """The failure this module exists to avoid: a plain openpyxl load-and-save loses the text box."""
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    wb["Cash Flow"]["B9"] = 1000
    path = tmp_path / "resaved.xlsx"
    wb.save(path)
    assert DISCLAIMER not in part(path.read_bytes(), "xl/drawings/drawing1.xml").decode()
    assert DISCLAIMER in part(template_bytes, "xl/drawings/drawing1.xml").decode()
